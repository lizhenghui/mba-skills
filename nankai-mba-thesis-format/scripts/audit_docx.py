#!/usr/bin/env python3
"""Read-only structural audit for Nankai University MBA thesis DOCX files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


CHINESE_NUM = "一二三四五六七八九十百零〇"
CHAPTER_RE = re.compile(rf"^第([{CHINESE_NUM}]+)章[\s\u3000]+(.+)$")
SECTION_RE = re.compile(rf"^第([{CHINESE_NUM}]+)节[\s\u3000]+(.+)$")
ITEM_RE = re.compile(rf"^([{CHINESE_NUM}]+)、\s*(.+)$")
SUBITEM_RE = re.compile(rf"^（([{CHINESE_NUM}]+)）\s*(.+)$")
FIGURE_RE = re.compile(r"^图(\d+)(?:\.(\d+))?[\s\u3000]+(.+)$")
TABLE_RE = re.compile(r"^表(\d+)(?:\.(\d+))?[\s\u3000]+(.+)$")
SOURCE_RE = re.compile(r"^(?:资料来源|数据来源|来源)\s*[：:]")
NUMERIC_CITATION_RE = re.compile(r"\[(?:\d+)(?:\s*[-–—,，]\s*\d+)*\]")
TITLE_PUNCT_RE = re.compile(r"[，。；：！？、,:;!?]")


def cm(value) -> float | None:
    return None if value is None else round(value.cm, 3)


def pt(value) -> float | None:
    return None if value is None else round(value.pt, 2)


def iter_blocks(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def next_nonempty_paragraph(paragraphs, index, direction=1):
    i = index + direction
    while 0 <= i < len(paragraphs):
        if paragraphs[i].text.strip() or paragraphs[i]._p.xpath(".//w:drawing | .//w:pict"):
            return paragraphs[i]
        i += direction
    return None


def normalized_label(text):
    return re.sub(r"[\s\u3000]+", "", text or "")


def effective_keep_with_next(paragraph):
    direct = paragraph.paragraph_format.keep_with_next
    if direct is not None:
        return direct
    style = paragraph.style
    while style is not None:
        value = style.paragraph_format.keep_with_next
        if value is not None:
            return value
        style = style.base_style
    return False


def border_value(container, edge):
    if container is None:
        return None
    node = container.find(qn(f"w:{edge}"))
    return None if node is None else node.get(qn("w:val"))


def table_border(table, edge):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    value = border_value(borders, edge)
    if value not in (None, "nil", "none"):
        return value
    rows = table.rows
    if not rows:
        return value
    row = rows[0] if edge in ("top", "insideH") else rows[-1]
    cell_edge = "bottom" if edge in ("bottom", "insideH") else edge
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.find(qn("w:tcBorders"))
        cell_value = border_value(tc_borders, cell_edge)
        if cell_value not in (None, "nil", "none"):
            return cell_value
    return value


def audit(path: Path):
    findings = []

    def add(severity, code, message, location=None):
        item = {"severity": severity, "code": code, "message": message}
        if location:
            item["location"] = location
        findings.append(item)

    try:
        with ZipFile(path) as package:
            bad = package.testzip()
            if bad:
                add("error", "DOCX_PACKAGE", f"DOCX包损坏：{bad}")
    except BadZipFile:
        return {"file": str(path), "findings": [{"severity": "error", "code": "DOCX_PACKAGE", "message": "不是有效的DOCX文件"}]}

    doc = Document(path)
    paragraphs = doc.paragraphs

    # Page geometry.
    expected = {"width": 21.0, "height": 29.7, "top": 3.8, "bottom": 3.8,
                "left": 3.2, "right": 3.2, "header": 3.0, "footer": 3.0}
    for index, section in enumerate(doc.sections, 1):
        actual = {"width": cm(section.page_width), "height": cm(section.page_height),
                  "top": cm(section.top_margin), "bottom": cm(section.bottom_margin),
                  "left": cm(section.left_margin), "right": cm(section.right_margin),
                  "header": cm(section.header_distance), "footer": cm(section.footer_distance)}
        for key, target in expected.items():
            if actual[key] is None or abs(actual[key] - target) > 0.08:
                add("error", "PAGE_GEOMETRY", f"{key}应为{target:.1f}cm，实为{actual[key]}cm", f"第{index}节")

    # Header underline (solid line below header text).
    for index, section in enumerate(doc.sections, 1):
        try:
            header = section.header
            header_paras = [p for p in header.paragraphs if p.text.strip()]
        except Exception:
            continue
        if not header_paras:
            continue
        underlined = False
        for p in header_paras:
            p_pr = p._p.find(qn("w:pPr"))
            if p_pr is None:
                continue
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is None:
                continue
            bottom = p_bdr.find(qn("w:bottom"))
            if bottom is not None and bottom.get(qn("w:val")) not in (None, "nil", "none"):
                underlined = True
                break
        if not underlined:
            add("warning", "HEADER_UNDERLINE", f"第{index}节页眉文字未检测到下划实线", f"第{index}节页眉")

    # Key style diagnostics.
    style_specs = {
        "Normal": (12, WD_ALIGN_PARAGRAPH.JUSTIFY),
        "Heading 1": (16, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (14, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 3": (13, None),
        "Heading 4": (12, None),
    }
    for name, (size, alignment) in style_specs.items():
        try:
            style = doc.styles[name]
        except KeyError:
            add("warning", "STYLE_MISSING", f"缺少命名样式 {name}")
            continue
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        if style.font.size is not None and abs(style.font.size.pt - size) > 0.2:
            add("warning", "STYLE_SIZE", f"{name}字号应为{size}磅，实为{pt(style.font.size)}磅")
        if alignment is not None and style.paragraph_format.alignment not in (None, alignment):
            add("warning", "STYLE_ALIGN", f"{name}对齐方式不符合MBA模式二")
    normal = doc.styles["Normal"]
    pf = normal.paragraph_format
    if pf.line_spacing is not None and hasattr(pf.line_spacing, "pt"):
        if abs(pf.line_spacing.pt - 20) > 0.2:
            add("warning", "BODY_LINE_SPACING", f"正文固定行距应为20磅，样式中为{pf.line_spacing.pt:.1f}磅")
    if pf.line_spacing_rule not in (None, WD_LINE_SPACING.EXACTLY):
        add("warning", "BODY_LINE_SPACING", "正文应使用固定值20磅行距")

    # Semantic heading hierarchy.
    chapters = []
    current_chapter = None
    current_section = None
    heading_counts = Counter()
    body_start_index = next((i for i, p in enumerate(paragraphs)
                             if CHAPTER_RE.match(p.text.strip()) and p.style.name == "Heading 1"), None)
    if body_start_index is None:
        chapter_occurrences = [i for i, p in enumerate(paragraphs) if CHAPTER_RE.match(p.text.strip())]
        body_start_index = chapter_occurrences[-1] if chapter_occurrences else 0
    for index, paragraph in enumerate(paragraphs[body_start_index:], body_start_index):
        text = paragraph.text.strip()
        match = CHAPTER_RE.match(text)
        level = None
        title = None
        if match:
            level, title = 1, match.group(2)
            current_chapter = {"title": text, "sections": []}
            chapters.append(current_chapter)
            current_section = None
        elif SECTION_RE.match(text):
            level, title = 2, SECTION_RE.match(text).group(2)
            if current_chapter is None:
                add("error", "HEADING_ORDER", "节标题出现在第一章之前", f"段落{index + 1}")
            else:
                current_section = {"title": text, "items": 0}
                current_chapter["sections"].append(current_section)
        elif ITEM_RE.match(text):
            level, title = 3, ITEM_RE.match(text).group(2)
            if current_section is None:
                add("error", "HEADING_ORDER", "目标题未归入任何节", f"段落{index + 1}")
            else:
                current_section["items"] += 1
        elif SUBITEM_RE.match(text):
            level, title = 4, SUBITEM_RE.match(text).group(2)

        if level:
            heading_counts[f"level_{level}"] += 1
            expected_style = f"Heading {level}"
            if level <= 3 and paragraph.style.name != expected_style:
                add("warning", "HEADING_STYLE", f"{text} 应使用 {expected_style} 或等效语义样式", f"段落{index + 1}")
            if title and TITLE_PUNCT_RE.search(title):
                add("error", "HEADING_PUNCTUATION", f"标题文字含标点：{text}", f"段落{index + 1}")
            if not effective_keep_with_next(paragraph):
                add("warning", "HEADING_PAGE_END", f"标题未设置与下段同页，可能落在页末：{text}", f"段落{index + 1}")

    for chapter in chapters:
        if len(chapter["sections"]) < 2:
            add("error", "CHAPTER_SECTIONS", f"{chapter['title']} 少于两节")
        for section in chapter["sections"]:
            if section["items"] == 0:
                add("warning", "SECTION_ITEMS", f"{section['title']} 下没有“目”级标题")

    # Abstracts and keywords.
    def find_exact(value):
        target = normalized_label(value)
        return next((i for i, p in enumerate(paragraphs) if normalized_label(p.text.strip()) == target), None)

    zh_idx, en_idx, toc_idx = find_exact("摘要"), find_exact("Abstract"), find_exact("目录")
    if zh_idx is None:
        add("error", "ABSTRACT_MISSING", "缺少中文摘要标题")
    if en_idx is None:
        add("error", "ABSTRACT_MISSING", "缺少Abstract标题")
    if zh_idx is not None and en_idx is not None:
        zh_text = "".join(p.text for p in paragraphs[zh_idx + 1:en_idx] if not p.text.strip().startswith("关键词"))
        zh_chars = len(re.findall(r"[\u3400-\u9fff]", zh_text))
        if not 300 <= zh_chars <= 1000:
            add("warning", "ABSTRACT_LENGTH", f"中文摘要约{zh_chars}个汉字，规范建议300–1000字")
    if en_idx is not None:
        end = toc_idx if toc_idx is not None and toc_idx > en_idx else min(len(paragraphs), en_idx + 30)
        en_text = " ".join(p.text for p in paragraphs[en_idx + 1:end] if not re.match(r"^Key\s*Words", p.text.strip(), re.I))
        en_words = len(re.findall(r"\b[A-Za-z]+(?:[-'][A-Za-z]+)*\b", en_text))
        if en_words < 300:
            add("warning", "ABSTRACT_LENGTH", f"英文摘要约{en_words}个英文实词，规范建议不少于300词")

    keyword_lines = [(i, p.text.strip()) for i, p in enumerate(paragraphs)
                     if re.match(r"^(关键词|Key\s*Words)\s*[：:]", p.text.strip(), re.I)]
    for index, text in keyword_lines:
        payload = re.split(r"[：:]", text, maxsplit=1)[-1].strip().rstrip("；;")
        terms = [part.strip() for part in re.split(r"[；;]", payload) if part.strip()]
        if not 3 <= len(terms) <= 8:
            add("error", "KEYWORD_COUNT", f"关键词应为3–8个，实为{len(terms)}个", f"段落{index + 1}")
        if text.startswith("关键词") and ";" in payload:
            add("warning", "KEYWORD_SEPARATOR", "中文论文关键词建议统一使用中文分号", f"段落{index + 1}")
    if len(keyword_lines) < 2:
        add("error", "KEYWORDS_MISSING", "中文或英文关键词行缺失")

    # TOC checks.
    field_parts = [doc.element]
    for section in doc.sections:
        field_parts.extend([section.header._element, section.footer._element])
    xml_text = " ".join(text for part in field_parts for text in part.xpath(".//w:instrText/text()"))
    has_toc_field = "TOC" in xml_text.upper()
    if toc_idx is None:
        add("error", "TOC_MISSING", "缺少目录标题")
    elif body_start_index is not None:
        toc_visible = [p.text.strip() for p in paragraphs[toc_idx + 1:body_start_index] if p.text.strip()]
        if toc_visible and not CHAPTER_RE.match(toc_visible[0]):
            add("error", "TOC_START", f"目录应直接从第一章开始，当前首项为：{toc_visible[0][:40]}")
        if toc_visible and not any(ITEM_RE.match(t) for t in toc_visible):
            add("warning", "TOC_LEVELS", "目录可见文本未体现“目”级，MBA目录应有三级")
        if not toc_visible and not has_toc_field:
            add("error", "TOC_EMPTY", "目录无可见条目且未检测到TOC域")

    # Figures.
    figure_captions = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        match = FIGURE_RE.match(text)
        if not match:
            continue
        figure_captions.append((int(match.group(1)), int(match.group(2) or 0), text))
        previous = next_nonempty_paragraph(paragraphs, index, -1)
        if previous is None or not previous._p.xpath(".//w:drawing | .//w:pict"):
            add("warning", "FIGURE_CAPTION_POSITION", f"图题应紧接在图下方：{text}", f"段落{index + 1}")
        following = next_nonempty_paragraph(paragraphs, index, 1)
        if following is None or not SOURCE_RE.match(following.text.strip()):
            add("error", "FIGURE_SOURCE", f"图后缺少资料来源：{text}", f"段落{index + 1}")
    if len(doc.inline_shapes) != len(figure_captions):
        add("warning", "FIGURE_COUNT", f"内嵌图形{len(doc.inline_shapes)}个，图题{len(figure_captions)}个，请人工核对浮动图形和图题")
    by_chapter = {}
    for chapter, number, text in figure_captions:
        by_chapter.setdefault(chapter, []).append((number, text))
    for chapter, items in by_chapter.items():
        nums = [number for number, _ in items]
        if 0 not in nums and nums != list(range(1, len(nums) + 1)):
            add("error", "FIGURE_NUMBERING", f"第{chapter}章图号不连续：{nums}")

    # Tables and their neighboring paragraphs.
    blocks = list(iter_blocks(doc))
    # 前置部分的封面/题名页/声明页表格属于学校固定版式，不参与表题、来源和三线表检查
    reached_body = body_start_index in (0, None)
    body_start_el = None if reached_body else paragraphs[body_start_index]._p
    table_index = 0
    for index, block in enumerate(blocks):
        if isinstance(block, Paragraph) and block._p is body_start_el:
            reached_body = True
        if not isinstance(block, Table):
            continue
        if not reached_body:
            continue
        table_index += 1
        prev_text = ""
        next_text = ""
        for prior in reversed(blocks[:index]):
            if isinstance(prior, Paragraph) and prior.text.strip():
                prev_text = prior.text.strip()
                break
        for following in blocks[index + 1:]:
            if isinstance(following, Paragraph) and following.text.strip():
                next_text = following.text.strip()
                break
        if not TABLE_RE.match(prev_text):
            add("error", "TABLE_CAPTION_POSITION", f"第{table_index}个表格上方缺少规范表题")
        if not SOURCE_RE.match(next_text):
            add("error", "TABLE_SOURCE", f"第{table_index}个表格后缺少资料来源")

        top = table_border(block, "top")
        bottom = table_border(block, "bottom")
        inside = table_border(block, "insideH")
        tbl_borders = block._tbl.tblPr.find(qn("w:tblBorders"))
        forbidden = {edge: border_value(tbl_borders, edge) for edge in ("left", "right", "insideV")}
        if top in (None, "nil", "none") or bottom in (None, "nil", "none") or inside in (None, "nil", "none"):
            add("warning", "THREE_LINE_TABLE", f"第{table_index}个表格未检测到完整三线表边框")
        if any(value not in (None, "nil", "none") for value in forbidden.values()):
            add("error", "THREE_LINE_TABLE", f"第{table_index}个表格存在侧边或内部竖线")
        if block.rows:
            tr_pr = block.rows[0]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                add("warning", "TABLE_HEADER_REPEAT", f"第{table_index}个表格首行未设置跨页重复表头")

    # Citations and references.
    ref_idx = find_exact("参考文献")
    if ref_idx is None:
        add("error", "REFERENCES_MISSING", "缺少参考文献部分")
        refs = []
    else:
        end = len(paragraphs)
        for i in range(ref_idx + 1, len(paragraphs)):
            text = normalized_label(paragraphs[i].text.strip())
            if text in {"致谢", "个人简历在学期间发表的学术论文与研究成果"} or text.startswith("附录"):
                end = i
                break
        refs = [p.text.strip() for p in paragraphs[ref_idx + 1:end] if p.text.strip()]
        body_text = "\n".join(p.text for p in paragraphs[:ref_idx])
        numeric = NUMERIC_CITATION_RE.findall(body_text)
        if numeric:
            add("error", "CITATION_SYSTEM", f"正文仍含顺序编码引文，例如：{', '.join(numeric[:5])}")
        if len(refs) < 30:
            add("warning", "REFERENCE_COUNT", f"参考文献{len(refs)}篇，MBA要求不宜少于30篇")
        english = [ref for ref in refs if re.match(r"^[A-Za-z]", re.sub(r"^\[\d+\]\s*", "", ref))]
        if len(english) < 5:
            add("warning", "ENGLISH_REFERENCE_COUNT", f"英文参考文献{len(english)}篇，MBA要求不宜少于5篇")
        if not all(re.match(r"^\[\d+\]", ref) for ref in refs):
            add("warning", "REFERENCE_NUMBERING",
                "文后参考文献条目应统一加方括号顺序编码，形如[1]、[2]……")
        else:
            nums = [int(re.match(r"^\[(\d+)\]", ref).group(1)) for ref in refs]
            if nums != list(range(1, len(refs) + 1)):
                add("warning", "REFERENCE_NUMBERING_ORDER",
                    f"参考文献编号未从1连续编号：{[n for n in nums[:10]]}…共{len(nums)}条")
        seen_english = False
        for ref in refs:
            is_english = bool(re.match(r"^[A-Za-z]", re.sub(r"^\[\d+\]\s*", "", ref)))
            seen_english = seen_english or is_english
            if seen_english and not is_english:
                add("error", "REFERENCE_LANGUAGE_ORDER", "中文参考文献出现在外文参考文献之后")
                break
        theses = [ref for ref in refs if re.search(r"\[D(?:/[^\]]+)?\]", ref, re.I)]
        if theses:
            add("warning", "THESIS_REFERENCE", f"检测到{len(theses)}篇学位论文文献，MBA一般不引用学位论文")

    # Section page-number formats and fields.
    section_properties = doc.element.body.xpath(".//w:sectPr")
    page_formats = []
    for sect_pr in section_properties:
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is not None:
            page_formats.append({"format": pg_num.get(qn("w:fmt")), "start": pg_num.get(qn("w:start"))})
    field_counts = Counter(text.strip().split()[0].upper() for part in field_parts
                           for text in part.xpath(".//w:instrText/text()") if text.strip())
    if field_counts.get("PAGE", 0) == 0:
        add("error", "PAGE_FIELD", "未检测到PAGE页码域")
    if not any(item.get("format") == "upperRoman" for item in page_formats):
        add("error", "FRONT_PAGE_NUMBER", "未检测到前置部分大写罗马页码格式 upperRoman")
    if not any(item.get("format") == "decimal" and item.get("start") == "1" for item in page_formats):
        add("error", "BODY_PAGE_NUMBER", "未检测到正文从1开始的阿拉伯页码设置")

    manual_checks = [
        "渲染后确认中文摘要不超过1页、英文摘要以1页为宜且不超过2页，并核对中英文语义对应。",
        "逐页确认每节篇幅长于1页、除章尾外页面排满、标题不落在页末最后一行。",
        "确认所有图片打印清晰、无水印、图中文字小于正文；地图内容另做合规审查。",
        "核对文内著者—出版年引文与文后条目一一对应，并人工确认中文拼音/外文字母排序。",
        "确认前置部分大写罗马页码、正文阿拉伯页码，以及第一章和各章右页起排。",
        "提交前确认双面打印、90g白色A4纸、书脊和线装/热胶装订要求。",
    ]

    counts = Counter(item["severity"] for item in findings)
    return {
        "file": str(path.resolve()),
        "summary": {
            "paragraphs": len(paragraphs),
            "sections": len(doc.sections),
            "tables": len(doc.tables),
            "inline_shapes": len(doc.inline_shapes),
            "chapters": len(chapters),
            "headings": dict(heading_counts),
            "references": len(refs),
            "fields": dict(field_counts),
            "page_number_formats": page_formats,
            "findings": dict(counts),
        },
        "findings": findings,
        "manual_checks": manual_checks,
    }


def main():
    parser = argparse.ArgumentParser(description="Audit a DOCX against Nankai University MBA thesis-format rules.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    parser.add_argument("--strict", action="store_true", help="Return exit code 2 when error findings exist.")
    args = parser.parse_args()
    if not args.docx.exists():
        parser.error(f"file not found: {args.docx}")

    report = audit(args.docx)
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    for item in report["findings"]:
        location = f" ({item['location']})" if item.get("location") else ""
        print(f"[{item['severity'].upper()}] {item['code']}{location}: {item['message']}")
    print("\nMANUAL CHECKS")
    for item in report["manual_checks"]:
        print(f"- {item}")

    if args.strict and any(item["severity"] == "error" for item in report["findings"]):
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
